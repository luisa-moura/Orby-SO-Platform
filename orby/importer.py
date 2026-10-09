from __future__ import annotations

import csv
from pathlib import Path

def read_table(path: str) -> tuple[list[str], list[dict[str, str]]]:
    file_path = Path(path)
    if file_path.suffix.lower() == ".csv":
        with file_path.open("r", encoding="utf-8-sig", newline="") as file:
            reader = csv.DictReader(file)
            headers = reader.fieldnames or []
            return headers, [{key: (value or "").strip() for key, value in row.items()} for row in reader]
    from openpyxl import load_workbook
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    sheet = workbook.active
    values = list(sheet.iter_rows(values_only=True))
    if not values:
        return [], []
    headers = [str(value).strip() if value is not None else "" for value in values[0]]
    rows = []
    for data in values[1:]:
        if any(value not in (None, "") for value in data):
            rows.append({headers[index]: str(value).strip() if value is not None else "" for index, value in enumerate(data)})
    return headers, rows


def validate_customers(rows: list[dict[str, str]], mapping: dict[str, str], duplicate_check) -> tuple[list[dict], list[dict]]:
    approved, rejected, seen = [], [], set()
    for index, source in enumerate(rows, start=2):
        item = {field: source.get(column, "").strip() for field, column in mapping.items() if column}
        errors = []
        if not item.get("nome"):
            errors.append("Nome obrigatório")
        if not item.get("telefone"):
            errors.append("Telefone obrigatório")
        key = (item.get("nome", "").casefold(), item.get("telefone", ""))
        if key in seen or (item.get("nome") and item.get("telefone") and duplicate_check(item["nome"], item["telefone"])):
            errors.append("Possível duplicidade")
        seen.add(key)
        if errors:
            rejected.append({"linha": index, "dados": item, "motivo": "; ".join(errors)})
        else:
            approved.append(item)
    return approved, rejected
